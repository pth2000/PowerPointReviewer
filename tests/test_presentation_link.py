"""验证路径迁移、讲稿保留和工程包跨设备恢复。"""

import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from pptx import Presentation

from app import paths, presentation_link, project_package


def make_ppt(path, count=2):
    ppt = Presentation()
    for index in range(count):
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = f'原始备注 {index + 1}'
    ppt.save(path)
    return Path(path)


class PresentationLinkTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)
        self.source = make_ppt(self.root / '原文稿.pptx')

    def test_renamed_or_moved_file_matches_by_content(self):
        info = presentation_link.describe(self.source)
        moved = self.root / '新名称.pptx'
        self.source.rename(moved)
        self.assertEqual(presentation_link.resolve(info, [moved]), moved)

    def test_same_name_and_page_count_do_not_prove_identity(self):
        info = presentation_link.describe(self.source)
        self.source.unlink()
        candidate = self.root / 'other' / self.source.name
        candidate.parent.mkdir()
        ppt = Presentation()
        for index in range(2):
            ppt.slides.add_slide(ppt.slide_layouts[6])
        ppt.save(candidate)
        self.assertIsNone(presentation_link.resolve(info, [candidate]))
        self.assertIsNone(presentation_link.resolve({'name': candidate.name}, [candidate]))

    def test_old_record_uses_original_path_without_requiring_fingerprint(self):
        info = presentation_link.from_record({'source_file': str(self.source)})
        self.assertEqual(presentation_link.resolve(info), self.source)
        self.assertEqual(presentation_link.from_record({'source_file': 'script.docx'}), {})

    def test_selected_ppt_must_cover_script_pages_and_remain_unchanged(self):
        original = self.source.read_bytes()
        with self.assertRaisesRegex(ValueError, '只有 2 页'):
            presentation_link.validate(self.source, [1, 3])
        info = presentation_link.validate(self.source, [2])
        self.assertEqual(info['slide_count'], 2)
        self.assertEqual(original, self.source.read_bytes())

    def test_link_update_preserves_original_script_audio_and_voice_profile(self):
        target = self.root / 'record.json'
        record = {'source_file': 'deleted.pptx', 'notes': {'1': 'AI 修改后的讲稿'},
                  'items': [{'text': '配音内容', 'cache_key': 'original'}],
                  'generation_profile': {'engine': 'original-voice'}, 'created_at': 'original-time'}
        target.write_text(json.dumps(record), encoding='utf-8')
        link = presentation_link.describe(self.source)
        presentation_link.update_record(target, link)
        updated = json.loads(target.read_text(encoding='utf-8'))
        self.assertEqual(updated.pop('presentation'), link)
        self.assertEqual(updated, record)
        self.assertFalse(list(self.root.glob('*.tmp')))


class PortablePresentationTests(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)
        self.source = make_ppt(self.root / '课程.pptx')
        self.audio = self.root / 'speech.wav'
        self.audio.write_bytes(b'generated audio')
        self.items = [{'page': 1, 'index': 0, 'text': '配音', 'duration': 1,
                       'cache_key': 'sample-key', 'cache_ext': 'wav', 'media': str(self.audio)}]
        for name, subfolder in [('AUDIO_CACHE_DIR', 'audio'), ('SESSION_DIR', 'sessions'),
                                ('PRESENTATION_DIR', 'presentations')]:
            destination = self.root / 'receiver' / subfolder
            destination.mkdir(parents=True)
            patcher = patch.object(paths, name, destination)
            patcher.start()
            self.addCleanup(patcher.stop)
        patcher = patch.object(paths, 'ensure_runtime_dirs')
        patcher.start()
        self.addCleanup(patcher.stop)

    def export(self, include_ppt=False):
        target = self.root / 'course.pprpkg'
        project_package.export_package(
            target, items=self.items, notes={1: '配音'}, mark='●', source_name='课程',
            generation_profile={'original': 'voice'}, speaker='speaker',
            presentation_path=self.source)
        if include_ppt:
            # 构造旧格式夹具，生产导出入口不再提供嵌入 PPT 功能。
            self.rewrite_archive(target, entry='presentation/document.pptx',
                                 payload=self.source.read_bytes())
        return target

    def record(self, result):
        return json.loads(result['record_path'].read_text(encoding='utf-8'))

    def rewrite_archive(self, target, *, entry=None, payload=None):
        with zipfile.ZipFile(target) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        manifest = json.loads(files['manifest.json'])
        if entry:
            manifest['presentation']['entry'] = entry
        if payload:
            files['presentation/document.pptx'] = payload
        files['manifest.json'] = json.dumps(manifest).encode('utf-8')
        with zipfile.ZipFile(target, 'w') as archive:
            for name, data in files.items():
                archive.writestr(name, data)

    def test_package_with_ppt_restores_on_other_device_without_original_file(self):
        original = self.source.read_bytes()
        target = self.export(True)
        manifest = project_package.read_manifest(target)
        self.assertNotIn('path', manifest['presentation'])
        self.source.unlink()
        result = project_package.import_package(target)
        record = self.record(result)
        restored = Path(record['presentation']['path'])
        self.assertTrue(restored.is_relative_to(paths.PRESENTATION_DIR))
        self.assertEqual(restored.read_bytes(), original)
        self.assertEqual(presentation_link.resolve(record['presentation']), restored)
        self.assertEqual((paths.AUDIO_CACHE_DIR / 'sample-key.wav').read_bytes(), b'generated audio')
        self.assertEqual(record['generation_profile'], {'original': 'voice'})
        self.assertEqual(record['source_file'], '')
        modified_at = restored.stat().st_mtime_ns
        again = project_package.import_package(target)
        self.assertNotEqual(again['record_path'], result['record_path'])
        self.assertEqual(again['skipped'], 1)
        self.assertEqual(Path(self.record(again)['presentation']['path']), restored)
        self.assertEqual(restored.stat().st_mtime_ns, modified_at)

    def test_package_without_ppt_finds_identical_companion_with_different_name(self):
        target = self.export(False)
        moved = self.root / '已改名.pptx'
        self.source.rename(moved)
        record = self.record(project_package.import_package(target))
        self.assertEqual(Path(record['presentation']['path']), moved)

    def test_package_without_corresponding_ppt_still_restores_audio(self):
        target = self.export(False)
        with zipfile.ZipFile(target) as archive:
            self.assertEqual(set(archive.namelist()), {'manifest.json', 'audio/sample-key.wav'})
        manifest = project_package.read_manifest(target)
        self.assertNotIn('entry', manifest['presentation'])
        self.assertNotIn('path', manifest['presentation'])
        self.source.unlink()
        record = self.record(project_package.import_package(target))
        self.assertNotIn('path', record['presentation'])
        self.assertEqual(record['notes'], {'1': '配音'})
        self.assertEqual(len(record['items']), 1)

    def test_old_package_remains_importable(self):
        target = self.root / 'old.pprpkg'
        project_package.export_package(target, items=self.items, notes={1: '配音'}, mark='●',
                                       source_name='旧工程', generation_profile={}, speaker='old')
        with zipfile.ZipFile(target) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        manifest = json.loads(files['manifest.json'])
        manifest.pop('presentation')
        files['manifest.json'] = json.dumps(manifest).encode('utf-8')
        with zipfile.ZipFile(target, 'w') as archive:
            for name, data in files.items():
                archive.writestr(name, data)
        record = self.record(project_package.import_package(target))
        self.assertEqual(record['presentation'], {})
        self.assertEqual(len(record['items']), 1)

    def test_corrupt_ppt_is_rejected_before_restoring_audio_and_cleans_temporary_file(self):
        target = self.export(True)
        self.rewrite_archive(target, payload=b'corrupt presentation')
        with self.assertRaisesRegex(ValueError, '校验未通过'):
            project_package.import_package(target)
        self.assertFalse(list(paths.AUDIO_CACHE_DIR.iterdir()))
        self.assertFalse(list(paths.PRESENTATION_DIR.rglob('*.tmp')))

    def test_untrusted_entry_cannot_extract_outside_presentation_storage(self):
        target = self.export(True)
        self.rewrite_archive(target, entry='../../outside.pptx')
        with self.assertRaises(ValueError):
            project_package.import_package(target)
        self.assertFalse(list(paths.PRESENTATION_DIR.iterdir()))
        self.assertFalse(list(paths.AUDIO_CACHE_DIR.iterdir()))


if __name__ == '__main__':
    unittest.main()
