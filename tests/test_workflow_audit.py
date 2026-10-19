"""覆盖跨操作顺序、失败恢复、导出状态和后台任务关闭的回归场景。"""

import json
from pathlib import Path
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace as NS
import unittest
from unittest.mock import Mock, patch

from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, QThread, QTimer
from qfluentwidgets import PushButton

from app import paths
from app.presentation import link as presentation_link
from app.documents import package as project_package
from ui import theme
from app.audio.playback import PlaybackBus
from app.settings_store import AppSettings
from app.window import Window
from tasks.audio_generation_task import AudioGenerationTask, GenerationResult
from tasks.preview_task import PreviewTask
from tts_engine import TTSEngine
from test_presentation_link import make_ppt
from ui.dialogs.rewrite_dialog import RewriteDialog
from ui.dialogs.session_history_dialog import SessionHistoryDialog
from ui.pages.reviewer_page import MissingAudioError, PPTReviewer
from ui.pages.settings_page import SettingInterface
from ui.pages.tools_page import ToolsInterface

from qt_test_case import QtTestCase


class WorkflowAuditTests(QtTestCase):
    def test_comment_exports_keep_raw_notes_and_ai_only_receives_body(self):
        from app.documents import io as script_io
        self.ctx.app_settings.set('script_comments_enabled', True)
        self.page.notes[1] = '// 教学提醒●\n第一页\n// 资料'
        self.page.mark_split()
        with patch('ui.pages.tools_page.RewriteDialog') as dialog:
            dialog.return_value.exec.return_value = False
            self.tools.open_rewrite_dialog()
        self.assertEqual(dialog.call_args.args[0], {1: '第一页', 3: '第三页'})
        with patch.object(self.tools, 'choose_export_dir', return_value=self.root), \
             patch.object(self.tools, 'get_ppt_path', return_value=self.source):
            self.tools.write_to_word()
            self.tools.write_to_json()
            self.tools.write_to_ppt()
        for suffix in ('docx', 'json'):
            exported = next(self.root.glob('*_Notes.' + suffix))
            self.assertEqual(script_io.load_script(exported).notes, self.page.notes)
        ppt = next(self.root.glob('*_NEW.pptx'))
        self.assertEqual(script_io.load_script(ppt).notes[1].strip(), self.page.notes[1])
        self.tools.create_error_info_bar.assert_not_called()

    def test_package_without_audio_regenerates_with_its_comment_rule(self):
        self.page.notes[1] = '// 教学提醒●\n第一页'
        package = self.root / 'comments.pprpkg'
        items = project_package.build_items(self.page.notes_list, self.page.notes_duration_list,
                                            self.page.note_cache_keys, self.page.note_cache_exts,
                                            self.page.media_list)
        project_package.export_package(package, items=items, notes=self.page.notes, mark='●',
                                       source_name='讲稿', generation_profile={}, speaker='',
                                       include_audio=False, script_comments_enabled=True)
        for audio in self.page.media_list:
            audio.unlink()
        self.assertFalse(self.ctx.app_settings.get('script_comments_enabled'))
        with patch.object(self.page.save_thread, 'start') as start, \
             patch.object(self.page.save_thread, 'configure') as configure, \
             patch.object(self.page, 'clean_temp_folder'):
            self.assertTrue(self.page.import_file(package))
        self.assertTrue(self.page.script_comments_enabled)
        self.assertEqual(configure.call_args.args[0],
                         [{'page': 1, 'text': '第一页'}, {'page': 3, 'text': '第三页'}])
        start.assert_called_once_with()


    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)
        for name, value in [('AUDIO_CACHE_DIR', 'audio'), ('SESSION_DIR', 'history'), ('TEMP_DIR', 'temp')]:
            target = self.root / value
            target.mkdir()
            patcher = patch.object(paths, name, target)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.engine = TTSEngine()
        self.engine.set_mode('edge')
        self.ctx = NS(version='1.6.3', updated_to='', tts_engine=self.engine, app_settings=AppSettings(),
                      config=Mock(), playback_bus=PlaybackBus(), hotkeys=Mock(), slideshow_watcher=Mock())
        self.ctx.config.save.return_value = True
        with patch('app.paths.ensure_runtime_dirs'):
            self.page = PPTReviewer(self.ctx)
        self.addCleanup(self.dispose)
        self.source = make_ppt(self.root / 'lesson.pptx', 3)
        self.page.note_file_path = self.page.presentation_path = self.source
        self.page.presentation_link = presentation_link.describe(self.source)
        self.page.note_file_name = 'lesson'
        self.page.notes = {1: '第一页', 3: '第三页'}
        self.page.mark_split()
        self.page.media_list = []
        for key in ('first', 'third'):
            audio = paths.AUDIO_CACHE_DIR / (key + '.wav')
            AudioGenerationTask._write_silence(audio, 'wav', seconds=0.1)
            self.page.media_list.append(audio)
        self.page.notes_duration_list = [0.1, 0.1]
        self.page.note_cache_keys = ['first', 'third']
        self.page.note_cache_exts = ['wav', 'wav']
        self.page.is_import = True
        self.page.check_import()
        self.tools = ToolsInterface(self.ctx, self.page)
        for target in (self.page, self.tools):
            for method in ('create_success_info_bar', 'create_warning_info_bar', 'create_error_info_bar'):
                patcher = patch.object(target, method)
                patcher.start()
                self.addCleanup(patcher.stop)

    def dispose(self):
        self.page.presentation_monitor.shutdown()
        self.tools.deleteLater()
        self.page.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        self.qt.processEvents()

    def make_settings(self):
        with patch('app.paths.ensure_runtime_dirs'), \
             patch.object(self.engine, 'get_voices_list', return_value=['zh-CN-XiaoxiaoNeural']), \
             patch('ui.theme.available_fonts', return_value=(('', theme.DEFAULT_FONT_LABEL),)):
            settings = SettingInterface(self.ctx)
        self.addCleanup(settings.deleteLater)
        return settings

    def record(self, items=None):
        path = paths.SESSION_DIR / 'record.json'
        profile = {'mode': 'bailian', 'options': {'model': 'cosyvoice-v3-flash', 'api_key': 'test-only-key'}}
        path.write_text(json.dumps({
            'source_file': str(self.source), 'source_name': 'lesson', 'notes': self.page.notes,
            'mark': '●', 'generation_profile': profile, 'speaker': '原发音人',
            'items': items if items is not None else [
                {'page': 1, 'text': '第一页', 'cache_key': 'first', 'cache_ext': 'wav', 'duration': 0.1},
                {'page': 3, 'text': '第三页', 'cache_key': 'third', 'cache_ext': 'wav', 'duration': 0.1},
            ]}), encoding='utf-8')
        return path

    def test_failed_history_load_preserves_working_session_and_playback(self):
        self.page.current_index = 1
        self.page._playing_sequence = True
        path = self.record([{'page': 1, 'text': '缺失配音', 'cache_key': 'missing'}])
        notes, media = dict(self.page.notes), list(self.page.media_list)
        with patch.object(self.page, 'stop_audio') as stop:
            with self.assertRaises(MissingAudioError):
                self.page.load_session_record(path)
        stop.assert_not_called()
        self.assertEqual(self.page.notes, notes)
        self.assertEqual(self.page.media_list, media)
        self.assertEqual(self.page.current_index, 1)
        self.assertTrue(self.page.is_import)
        self.assertTrue(self.page._playing_sequence)

    def test_unplayable_audio_restores_controls_and_does_not_advance(self):
        from PySide6.QtMultimedia import QMediaPlayer
        self.page._playing_sequence = True
        self.page.playButton.setEnabled(False)
        self.page.next_timer.start(100)
        with patch.object(self.page.player, 'error', return_value=QMediaPlayer.Error.FormatError):
            self.page.on_playback_error(QMediaPlayer.Error.FormatError, 'invalid audio')
        self.assertFalse(self.page._playing_sequence)
        self.assertFalse(self.page.next_timer.isActive())
        self.assertTrue(self.page.playButton.isEnabled())
        self.assertEqual(self.page.current_index, 0)
        self.page.create_error_info_bar.assert_called_once()

    def test_late_audio_error_after_stop_does_not_interrupt_other_playback(self):
        from PySide6.QtMultimedia import QMediaPlayer
        self.page._playing_sequence = False
        with patch.object(self.page, 'stop_audio') as stop:
            self.page.on_playback_error(QMediaPlayer.Error.ResourceError, 'late error')
        stop.assert_not_called()
        self.page.create_error_info_bar.assert_not_called()

    def test_old_audio_error_is_ignored_when_player_has_recovered_or_loaded_new_source(self):
        from PySide6.QtMultimedia import QMediaPlayer
        self.page._playing_sequence = True
        with patch.object(self.page, 'stop_audio') as stop:
            self.page.on_playback_error(QMediaPlayer.Error.ResourceError, 'stale error')
        stop.assert_not_called()
        self.page.create_error_info_bar.assert_not_called()

    def test_sparse_pages_can_jump_to_last_page_and_missing_page_does_not_stop(self):
        self.page.load_session_record(self.record())
        self.assertEqual(self.page.pageJumpSpinBox.maximum(), 3)
        self.page.pageJumpSpinBox.setValue(3)
        self.page.jump_page()
        self.assertEqual(self.page.current_index, 1)
        self.page.pageJumpSpinBox.setValue(2)
        with patch.object(self.page, 'stop_audio') as stop, patch('app.presentation.slideshow.goto') as goto:
            self.page.jump_page()
        stop.assert_not_called()
        goto.assert_not_called()
        self.assertEqual(self.page.current_index, 1)

    def test_ppt_note_export_rejects_removed_slides_instead_of_silently_omitting_script(self):
        make_ppt(self.source, 1)
        with patch.object(self.tools, 'choose_export_dir') as choose:
            self.tools.write_to_ppt()
        choose.assert_not_called()
        self.tools.create_warning_info_bar.assert_called_once()
        self.assertEqual(self.page.notes[3], '第三页')

    def test_ppt_note_export_preserves_original_and_writes_all_existing_script_pages(self):
        from pptx import Presentation
        original = self.source.read_bytes()
        with patch.object(self.tools, 'choose_export_dir', return_value=self.root):
            self.tools.write_to_ppt()
        exported = Presentation(str(self.root / 'lesson_NEW.pptx'))
        self.assertEqual(exported.slides[0].notes_slide.notes_text_frame.text, '第一页')
        self.assertEqual(exported.slides[2].notes_slide.notes_text_frame.text, '第三页')
        self.assertEqual(self.source.read_bytes(), original)

    def test_failed_generation_can_be_reimported_without_forcing_cache_bypass(self):
        self.page.thread_save_error('test failure')
        self.page.refresh_presentation_state()
        self.assertFalse(self.page.regenerate_action.isEnabled())
        self.assertEqual(self.page.regenerate_action.text(), '忽略缓存重新生成')
        self.assertTrue(self.page.getFileButton.isEnabled())
        with patch.object(self.page, 'regenerate') as generate:
            self.page.force_regenerate()
        generate.assert_not_called()
        with patch.object(self.page.save_thread, 'configure') as configure, \
             patch.object(self.page.save_thread, 'start') as start, \
             patch.object(self.page, 'clean_temp_folder'):
            self.assertTrue(self.page.import_file(self.source))
        start.assert_called_once_with()
        self.assertFalse(configure.call_args.kwargs['force_regenerate'])

    def test_record_save_failure_keeps_generated_audio_playable(self):
        result = GenerationResult(list(self.page.media_list), [0.1, 0.1], ['first', 'third'], ['wav', 'wav'],
                                  generation_profile={'mode': 'edge'}, speaker='原音色')
        with patch.object(self.page, 'save_session_record', side_effect=OSError('disk full')):
            self.page.thread_save_finish(result)
        self.assertTrue(self.page.is_import)
        self.assertTrue(self.page.playCardWidget.isEnabled())
        self.assertTrue(self.page.getFileButton.isEnabled())
        self.page.create_error_info_bar.assert_called_once()

    def test_saved_and_exported_metadata_describe_loaded_audio_not_current_settings(self):
        self.page.load_session_record(self.record())
        self.engine.set_mode('edge')
        profile, speaker, _ = self.tools.audio_metadata()
        self.assertEqual(profile['mode'], 'bailian')
        self.assertEqual(speaker, '原发音人')
        self.page.save_session_record()
        saved = json.loads(self.page.session_record_path.read_text(encoding='utf-8'))
        self.assertEqual(saved['generation_profile'], profile)
        self.assertEqual(saved['speaker'], speaker)

    def test_incomplete_or_busy_audio_disables_exports_and_rechecks_before_choosing_directory(self):
        self.page.media_list.pop()
        self.tools.refresh_state()
        audio_buttons = [button for card in self.tools.export_cards
                         for button in card.findChildren(PushButton)
                         if button.property('requires_audio')]
        self.assertTrue(audio_buttons)
        self.assertTrue(all(not button.isEnabled() for button in audio_buttons))
        with patch.object(self.tools, 'choose_export_dir') as choose:
            self.tools.export_audio_files()
            self.tools.export_merged_audio()
            self.tools.export_project_package()
        choose.assert_not_called()
        with patch.object(self.page, 'is_busy', return_value=True), \
             patch.object(self.tools, 'choose_export_dir') as choose:
            self.tools.write_to_word()
            self.tools.write_to_json()
        choose.assert_not_called()

    def test_repeated_audio_export_preserves_previous_directory(self):
        destination = self.root / 'export'
        destination.mkdir()
        with patch.object(self.tools, 'choose_export_dir', return_value=destination):
            self.tools.export_audio_files()
            previous = {path.name: path.read_bytes() for path in (destination / 'lesson_Audio').iterdir()}
            self.tools.export_audio_files()
        self.assertTrue((destination / 'lesson_Audio_2').is_dir())
        self.assertEqual(previous, {path.name: path.read_bytes() for path in (destination / 'lesson_Audio').iterdir()})

    def test_export_name_remains_inside_chosen_directory(self):
        self.page.note_file_name = '../课程:第一节/讲稿'
        name = self.tools.export_name()
        self.assertNotIn('/', name)
        self.assertNotIn(':', name)
        self.assertEqual((self.root / (name + '.json')).parent, self.root)

    def test_export_unique_path_does_not_overwrite_after_999_names(self):
        original = self.root / 'lecture.json'
        with patch.object(Path, 'exists', lambda path: path.name != 'lecture_1000.json'):
            self.assertEqual(self.tools.unique_path(original).name, 'lecture_1000.json')

    def test_partial_history_delete_keeps_audio_referenced_by_failed_record(self):
        for name in ('a', 'b'):
            (paths.SESSION_DIR / (name + '.json')).write_text(
                json.dumps({'items': [{'cache_key': 'shared'}]}), encoding='utf-8')
        audio = paths.AUDIO_CACHE_DIR / 'shared.wav'
        audio.write_bytes(b'keep shared audio')
        a, b = paths.SESSION_DIR / 'a.json', paths.SESSION_DIR / 'b.json'
        dialog = SessionHistoryDialog(paths.SESSION_DIR, self.page)
        self.addCleanup(dialog.deleteLater)
        unlink = Path.unlink
        def fail_one(path, *args, **kwargs):
            if path == b:
                raise PermissionError('record in use')
            return unlink(path, *args, **kwargs)
        with patch.object(dialog, 'selected_record_paths', return_value=[a, b]), \
             patch('ui.dialogs.session_history_dialog.DeleteRecordMessageBox') as confirm, \
             patch.object(Path, 'unlink', fail_one):
            confirm.return_value.exec.return_value = True
            confirm.return_value.should_delete_audio.return_value = True
            dialog.delete_selected()
        self.assertFalse(a.exists())
        self.assertTrue(b.exists())
        self.assertEqual(audio.read_bytes(), b'keep shared audio')

    def test_orphan_cleanup_preserves_audio_currently_loaded_without_a_record(self):
        orphan = paths.AUDIO_CACHE_DIR / 'orphan.wav'
        orphan.write_bytes(b'orphan')
        dialog = SessionHistoryDialog(paths.SESSION_DIR, self.page)
        self.addCleanup(dialog.deleteLater)
        dialog.clean_orphan_audio()
        self.assertFalse(orphan.exists())
        self.assertTrue(all(path.is_file() for path in self.page.media_list))

    def test_project_package_omits_credentials_without_changing_cache_identity(self):
        profile = {'mode': 'qwen_clone', 'options': {'api_key': 'test-only-key', 'voice': 'voice-1'}}
        items = project_package.build_items(self.page.notes_list, self.page.notes_duration_list,
                                             self.page.note_cache_keys, self.page.note_cache_exts, self.page.media_list)
        target = self.root / 'lecture.pprpkg'
        project_package.export_package(target, items=items, notes=self.page.notes, mark='●',
                                       source_name='lecture', generation_profile=profile, speaker='voice-1')
        manifest = project_package.read_manifest(target)
        self.assertNotIn('api_key', manifest['generation_profile']['options'])
        self.assertEqual(manifest['generation_profile']['options']['voice'], 'voice-1')
        self.assertEqual([item['cache_key'] for item in manifest['items']], ['first', 'third'])
        self.assertEqual(profile['options']['api_key'], 'test-only-key')

    def test_malformed_history_record_does_not_break_list_or_cache_cleanup(self):
        (paths.SESSION_DIR / 'invalid.json').write_text('[]', encoding='utf-8')
        dialog = SessionHistoryDialog(paths.SESSION_DIR, self.page)
        self.addCleanup(dialog.deleteLater)
        self.assertEqual(dialog.table.rowCount(), 1)
        dialog.clean_orphan_audio()
        self.assertTrue(all(path.is_file() for path in self.page.media_list))

    def test_voice_creation_does_not_select_it_and_cancelled_deletion_does_not_call_service(self):
        from ui.dialogs.qwen_clone_voice_dialog import QwenCloneVoiceDialog
        engine = TTSEngine()
        engine.set_mode('qwen_clone')
        engine.set_current_option('voice', 'original-voice')
        with patch('engines.qwen_clone.create_voice', return_value='new-voice'):
            self.assertEqual(engine.create_qwen_clone_voice('reference.wav'), 'new-voice')
        self.assertEqual(engine.get_current_option_values()['voice'], 'original-voice')
        with patch.object(QwenCloneVoiceDialog, 'refresh_list', lambda self, **kwargs: None):
            dialog = QwenCloneVoiceDialog(engine, self.page)
        self.addCleanup(dialog.deleteLater)
        with patch.object(dialog, '_current_row_voice', return_value='original-voice'), \
             patch('ui.dialogs.qwen_clone_voice_dialog.MessageBox') as confirm, \
             patch.object(engine, 'delete_qwen_clone_voice') as delete:
            confirm.return_value.exec.return_value = False
            dialog.delete_selected_voice()
        delete.assert_not_called()

    def test_package_without_audio_loads_script_and_prepares_audio_instead_of_failing(self):
        items = [{'page': 2, 'index': 0, 'text': '工程包讲稿', 'cache_key': 'missing', 'cache_ext': 'wav'}]
        target = self.root / 'no-audio.pprpkg'
        project_package.export_package(target, items=items, notes={2: '工程包讲稿'}, mark='●',
                                       source_name='工程包', generation_profile={}, speaker='', include_audio=False)
        with patch('app.paths.ensure_runtime_dirs'), patch.object(self.page, 'init_general_play') as generate:
            self.assertTrue(self.page.import_file(target))
        self.assertEqual(self.page.notes, {2: '工程包讲稿'})
        self.assertTrue(self.page.is_restricted())
        generate.assert_called_once()

    def test_changing_mark_regenerates_current_script_but_busy_change_is_blocked(self):
        with patch('ui.pages.reviewer_page.EditMarkMessageBox') as dialog, \
             patch.object(self.page, 'regenerate') as generate:
            dialog.return_value.exec.return_value = True
            dialog.return_value.urlLineEdit.text.return_value = '|'
            self.page.show_edit_mark_dialog()
        self.assertEqual(self.page.mark, '|')
        generate.assert_called_once_with()
        with patch.object(self.page, 'is_busy', return_value=True), \
             patch('ui.pages.reviewer_page.EditMarkMessageBox') as dialog:
            self.page.show_edit_mark_dialog()
        dialog.assert_not_called()

    def test_settings_save_failure_is_returned_to_the_caller(self):
        settings = self.make_settings()
        self.engine.set_current_option('rate', '+10%')
        self.ctx.config.save.return_value = False
        with patch('ui.pages.settings_page.MessageBox') as dialog, \
             patch.object(settings, 'create_error_info_bar'):
            dialog.return_value.exec.return_value = True
            self.assertFalse(settings.prompt_unsaved_changes(allow_regenerate=False))
        self.assertTrue(settings.is_dirty())

    def test_late_preview_does_not_interrupt_playback_started_in_another_page(self):
        settings = self.make_settings()
        settings._preview_requested = True
        settings.on_stop_requested(self.page)
        with patch.object(settings.preview_player, 'play') as play:
            settings.on_preview_generated(str(self.page.media_list[0]))
        play.assert_not_called()

    def test_drop_resolves_unsaved_settings_before_importing(self):
        settings = Mock()
        settings.is_dirty.return_value = True
        settings.prompt_unsaved_changes.return_value = False
        host = NS(setting_interface=settings, ppt_r=Mock(), switchTo=Mock())
        Window._import_dropped_files(host, ['D:/lecture.pptx'])
        host.ppt_r.import_file.assert_not_called()
        settings.prompt_unsaved_changes.return_value = True
        Window._import_dropped_files(host, ['D:/lecture.pptx'])
        host.ppt_r.import_file.assert_called_once_with('D:/lecture.pptx')

    def test_ai_close_requests_cancellation_without_blocking_or_closing_running_worker(self):
        dialog = RewriteDialog({1: '讲稿'}, '●', AppSettings())
        self.addCleanup(dialog.deleteLater)
        dialog.task = Mock()
        dialog.task.isRunning.return_value = True
        dialog.reject()
        dialog.task.cancel.assert_called_once()
        dialog.task.wait.assert_not_called()
        self.assertTrue(dialog._closing)
        self.assertFalse(dialog.close_button.isEnabled())
        dialog.task = None
        dialog._worker_finished()


class EngineSnapshotTests(QtTestCase):

    def test_snapshot_and_preview_keep_original_mode_and_parameters(self):
        engine = TTSEngine()
        engine.set_mode('bailian')
        engine.set_current_option('volume', 35)
        clone = engine.snapshot()
        preview = PreviewTask(engine, '试听', 'unused.wav')
        engine.set_current_option('volume', 80)
        engine.set_mode('edge')
        self.assertEqual(clone.get_mode(), 'bailian')
        self.assertEqual(clone.get_current_option_values()['volume'], 35)
        self.assertEqual(preview.tts_engine.get_mode(), 'bailian')
        self.assertEqual(preview.tts_engine.get_current_option_values()['volume'], 35)
        clone.set_current_option('volume', 45)
        self.assertEqual(engine._engine_settings['bailian']['volume'], 80)

    def test_generation_result_uses_configure_snapshot_after_ui_changes(self):
        engine = TTSEngine()
        engine.set_mode('bailian')
        engine.set_current_option('volume', 35)
        task = AudioGenerationTask(engine)
        task.configure([{'page': 1, 'text': '第一段'}, {'page': 1, 'text': '第二段'}], 0)
        engine.set_current_option('volume', 80)
        engine.set_mode('edge')
        with tempfile.TemporaryDirectory() as folder:
            task.audio_cache_path = Path(folder) / 'cache'
            task.output_path = Path(folder) / 'out'
            task.audio_cache_path.mkdir()
            task.output_path.mkdir()
            with patch.object(task.tts_engine, 'save_file',
                              side_effect=lambda text, path: AudioGenerationTask._write_silence(Path(path), 'wav')):
                result = task.save_wav()
        self.assertEqual(result.generation_profile['mode'], 'bailian')
        self.assertEqual(result.generation_profile['options']['volume'], 35)
        self.assertTrue(all(ext == 'wav' for ext in result.cache_exts))

    def test_parallel_bailian_initialization_keeps_each_account_key(self):
        import dashscope
        from engines import bailian
        observed = []
        class Callback:
            def __init__(self, path):
                self.path = Path(path)
                self.error_message = None
                self.done_event = threading.Event()
        class Synthesizer:
            def __init__(self, **kwargs):
                before = dashscope.api_key
                time.sleep(0.01)
                observed.append((before, dashscope.api_key))
                self.callback = kwargs['callback']
            def call(self, text):
                self.callback.path.write_bytes(b'audio')
                self.callback.done_event.set()
        with tempfile.TemporaryDirectory() as folder, \
             patch('engines.bailian._file_callback_class', return_value=Callback), \
             patch('dashscope.audio.tts_v2.SpeechSynthesizer', Synthesizer), \
             patch.object(dashscope, 'api_key', None):
            with ThreadPoolExecutor(max_workers=2) as pool:
                futures = [pool.submit(bailian.save, 'test', str(Path(folder) / f'{index}.wav'),
                                       api_key=key) for index, key in enumerate(('test-key-a', 'test-key-b'))]
                for future in futures:
                    future.result()
        self.assertEqual({before for before, after in observed}, {'test-key-a', 'test-key-b'})
        self.assertTrue(all(before == after for before, after in observed))


class ShutdownTests(QtTestCase):

    def test_window_waits_for_worker_and_then_closes_without_destroying_it(self):
        engine = TTSEngine()
        engine.set_mode('edge')
        ctx = NS(version='1.6.3', updated_to='', tts_engine=engine, app_settings=AppSettings(),
                 config=Mock(), playback_bus=PlaybackBus(), hotkeys=Mock(), slideshow_watcher=Mock())
        with patch('app.paths.ensure_runtime_dirs'), \
             patch.object(engine, 'get_voices_list', return_value=['zh-CN-XiaoxiaoNeural']), \
             patch('ui.theme.available_fonts', return_value=(('', theme.DEFAULT_FONT_LABEL),)):
            root = Window(ctx)
        release = threading.Event()
        started = threading.Event()
        class Worker(QThread):
            def run(self):
                started.set()
                release.wait(3)
        worker = Worker(root)
        def cleanup():
            release.set()
            worker.wait(3000)
            root.ppt_r.presentation_monitor.shutdown()
            root.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        self.addCleanup(cleanup)
        worker.start()
        self.assertTrue(started.wait(1))
        root.close()
        self.assertTrue(root.isVisible())
        self.assertTrue(root._waiting_to_close)
        self.assertFalse(root.isEnabled())
        release.set()
        loop = QEventLoop()
        worker.finished.connect(loop.quit)
        QTimer.singleShot(3000, loop.quit)
        loop.exec()
        self.qt.processEvents()
        self.assertFalse(worker.isRunning())
        self.assertFalse(root.isVisible())
        root.show()
        with patch.object(root.setting_interface, 'prompt_unsaved_changes', return_value=False):
            root.close()
        self.assertTrue(root.isVisible())


if __name__ == '__main__':
    unittest.main()
