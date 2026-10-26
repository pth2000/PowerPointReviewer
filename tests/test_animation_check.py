"""导入后的点击步骤预检与继续/取消生成流程。"""

from pathlib import Path
import tempfile
from types import SimpleNamespace as NS
import unittest
from unittest.mock import Mock, patch

from pptx import Presentation
from pptx.oxml import parse_xml
from PySide6.QtCore import QCoreApplication, QEvent, QTimer

from app.presentation.animation import AnimationIssue, check_alignment
from app.presentation.visibility import hidden_shape_ids
from app.audio.playback import PlaybackBus
from app.settings_store import AppSettings
from ui.dialogs.animation_check_dialog import AnimationCheckDialog

from qt_test_case import QtTestCase


P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def make_deck(path, pages):
    """Create ordinary click groups with optional automatic or interactive groups."""
    ppt = Presentation()
    for groups in pages:
        slide = ppt.slides.add_slide(ppt.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = '讲稿'
        if groups is None:
            continue
        nodes = []
        for index, effects in enumerate(groups):
            auto = all(kind in ('withEffect', 'afterEffect') for kind in effects)
            conditions = '<p:cond delay="indefinite"/>'
            if auto:
                conditions += '<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>'
            effects_xml = ''.join(f'<p:par><p:cTn nodeType="{kind}"/></p:par>' for kind in effects)
            nodes.append(f'<p:par><p:cTn id="{index + 3}"><p:stCondLst>{conditions}</p:stCondLst>'
                         f'<p:childTnLst>{effects_xml}</p:childTnLst></p:cTn></p:par>')
        slide._element.append(parse_xml(
            f'<p:timing xmlns:p="{P}"><p:tnLst><p:seq><p:cTn nodeType="mainSeq">'
            f'<p:childTnLst>{"".join(nodes)}</p:childTnLst></p:cTn></p:seq></p:tnLst></p:timing>'))
    ppt.save(path)
    return path


def segments(*counts):
    return [{'page': page, 'text': f'第 {part} 段'}
            for page, count in enumerate(counts, 1) for part in range(count)]


def media_timing_xml(identifier=7, *, interactive=False):
    player = (f'<p:video><p:cMediaNode mute="1"><p:cTn repeatCount="indefinite"/>'
              f'<p:tgtEl><p:spTgt spid="{identifier}"/></p:tgtEl></p:cMediaNode></p:video>')
    if interactive:
        player += (f'<p:seq><p:cTn nodeType="interactiveSeq"><p:stCondLst>'
                   f'<p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{identifier}"/>'
                   '</p:tgtEl></p:cond></p:stCondLst><p:childTnLst><p:par>'
                   '<p:cTn nodeType="clickEffect" presetClass="mediacall"><p:childTnLst>'
                   '<p:cmd type="call" cmd="togglePause"><p:cBhvr><p:cTn dur="1"/>'
                   f'<p:tgtEl><p:spTgt spid="{identifier}"/></p:tgtEl></p:cBhvr></p:cmd>'
                   '</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:seq>')
    return player


class AnimationCheckTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.source = Path(self.directory.name) / '讲稿.pptx'

    def test_click_groups_include_paragraphs_but_not_with_or_after_effects(self):
        make_deck(self.source, [None,
                               [('clickEffect', 'withEffect', 'afterEffect'), ('clickEffect',)],
                               [('clickEffect',), ('clickPar',), ('clickPar',)],
                               [('afterEffect', 'withEffect')]])
        original = self.source.read_bytes()
        self.assertEqual(check_alignment(self.source, segments(1, 3, 4, 1)), [])
        self.assertEqual(self.source.read_bytes(), original)

    def deck_with_hidden_targets(self, hidden, groups):
        """每个效果指定对象，隐藏标记与选择窗格保存的 XML 一致。"""
        make_deck(self.source, [[tuple(kind for kind, _target in group) for group in groups]])
        ppt = Presentation(self.source)
        slide = ppt.slides[0]
        tree = slide._element.find(f'{{{P}}}cSld/{{{P}}}spTree')
        for identifier in {target for group in groups for _kind, target in group}:
            tree.append(parse_xml(f'<p:sp xmlns:p="{P}"><p:nvSpPr>'
                f'<p:cNvPr id="{identifier}" name="对象{identifier}" hidden="{int(identifier in hidden)}"/>'
                '<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr/></p:sp>'))
        effects = [node for node in slide._element.iter(f'{{{P}}}cTn')
                   if node.get('nodeType') in ('clickEffect', 'clickPar', 'withEffect', 'afterEffect')]
        for node, (_kind, target) in zip(effects, [item for group in groups for item in group]):
            node.append(parse_xml(f'<p:cBhvr xmlns:p="{P}"><p:tgtEl>'
                                  f'<p:spTgt spid="{target}"/></p:tgtEl></p:cBhvr>'))
        ppt.save(self.source)

    def test_hidden_click_objects_do_not_require_extra_script_segments(self):
        self.deck_with_hidden_targets({3}, [[('clickEffect', 2)], [('clickEffect', 3)], [('clickEffect', 4)]])
        original = self.source.read_bytes()
        self.assertEqual(check_alignment(self.source, segments(3)), [])
        issues = check_alignment(self.source, segments(4))
        self.assertEqual((issues[0].clicks, issues[0].reason), (2, '讲稿多出 1 个点击步骤'))
        self.assertEqual(self.source.read_bytes(), original)

    def test_visible_with_effect_does_not_resurrect_hidden_click_leader(self):
        self.deck_with_hidden_targets({2}, [[('clickEffect', 2), ('withEffect', 3)], [('clickEffect', 4)]])
        self.assertEqual(check_alignment(self.source, segments(2)), [])

    def test_hidden_simultaneous_effect_does_not_remove_visible_click(self):
        self.deck_with_hidden_targets({3}, [[('clickEffect', 2), ('withEffect', 3)], [('clickEffect', 4)]])
        self.assertEqual(check_alignment(self.source, segments(3)), [])

    def test_all_hidden_clicks_and_hidden_paragraph_clicks_are_zero(self):
        self.deck_with_hidden_targets({2, 3}, [[('clickEffect', 2)], [('clickPar', 3)]])
        self.assertEqual(check_alignment(self.source, segments(1)), [])

    def test_hidden_group_hides_descendants_even_with_explicit_visible_flag(self):
        self.deck_with_hidden_targets(set(), [[('clickEffect', 2)], [('clickEffect', 3)]])
        ppt = Presentation(self.source)
        slide = ppt.slides[0]
        tree = slide._element.find(f'{{{P}}}cSld/{{{P}}}spTree')
        group = parse_xml(f'<p:grpSp xmlns:p="{P}"><p:nvGrpSpPr>'
            '<p:cNvPr id="10" name="隐藏组合" hidden="true"/><p:cNvGrpSpPr/><p:nvPr/>'
            '</p:nvGrpSpPr><p:grpSpPr/></p:grpSp>')
        for shape in list(tree):
            if shape.tag == f'{{{P}}}sp':
                tree.remove(shape)
                group.append(shape)
        tree.append(group)
        self.assertEqual(hidden_shape_ids(slide._element), {'10', '2', '3'})
        ppt.save(self.source)
        self.assertEqual(check_alignment(self.source, segments(1)), [])

    def test_objects_outside_page_are_still_counted_when_not_hidden(self):
        self.deck_with_hidden_targets(set(), [[('clickEffect', 2)]])
        ppt = Presentation(self.source)
        shape = ppt.slides[0]._element.find(f'{{{P}}}cSld/{{{P}}}spTree/{{{P}}}sp')
        shape.find(f'{{{P}}}spPr').append(parse_xml(
            '<a:xfrm xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            '<a:off x="-9999999" y="-9999999"/><a:ext cx="100" cy="100"/></a:xfrm>'))
        ppt.save(self.source)
        self.assertEqual(check_alignment(self.source, segments(2)), [])

    def test_only_mismatching_pages_are_reported_in_slide_order(self):
        make_deck(self.source, [[('clickEffect',)], [('clickEffect',), ('clickEffect',)], None])
        issues = check_alignment(self.source, segments(3, 2, 1))
        self.assertEqual([(i.page, i.segments, i.clicks) for i in issues], [(1, 3, 1), (2, 2, 2)])
        self.assertIn('多出 1', issues[0].reason)
        self.assertIn('还有 1', issues[1].reason)

    def test_interactive_media_and_unknown_animation_are_not_counted_as_zero(self):
        make_deck(self.source, [[('interactiveSeq',)], [('clickEffect',)], [('unknown',)]])
        ppt = Presentation(self.source)
        ppt.slides[1]._element.find(f'{{{P}}}timing').append(parse_xml(f'<p:video xmlns:p="{P}"/>'))
        ppt.save(self.source)
        issues = check_alignment(self.source, segments(1, 1, 1))
        self.assertEqual([i.page for i in issues], [1, 2, 3])
        self.assertTrue(all(i.clicks is None for i in issues))

    def media_deck(self, kind='withEffect', *, interactive=False):
        make_deck(self.source, [[(kind,)]])
        ppt = Presentation(self.source)
        timing = ppt.slides[0]._element.find(f'{{{P}}}timing')
        effect = next(node for node in timing.iter(f'{{{P}}}cTn') if node.get('nodeType') == kind)
        effect.set('presetClass', 'mediacall')
        effect.append(parse_xml(f'<p:childTnLst xmlns:p="{P}"><p:cmd type="call" cmd="playFrom(0.0)">'
                               '<p:cBhvr><p:cTn dur="2000"/><p:tgtEl><p:spTgt spid="7"/>'
                               '</p:tgtEl></p:cBhvr></p:cmd></p:childTnLst>'))
        nodes = parse_xml(f'<p:tnLst xmlns:p="{P}">{media_timing_xml(interactive=interactive)}</p:tnLst>')
        timing.find(f'{{{P}}}tnLst').extend(list(nodes))
        ppt.save(self.source)

    def test_autoplay_and_looping_video_do_not_warn_or_add_click_steps(self):
        self.media_deck(interactive=True)
        original = self.source.read_bytes()
        self.assertEqual(check_alignment(self.source, segments(1)), [])
        self.assertEqual(self.source.read_bytes(), original)

    def test_click_to_play_media_is_counted_as_one_sequential_step(self):
        self.media_deck('clickEffect', interactive=True)
        self.assertEqual(check_alignment(self.source, segments(2)), [])
        issue = check_alignment(self.source, segments(1))[0]
        self.assertEqual(issue.clicks, 1)

    def test_media_interaction_targeting_another_object_still_needs_confirmation(self):
        self.media_deck(interactive=True)
        ppt = Presentation(self.source)
        timing = ppt.slides[0]._element.find(f'{{{P}}}timing')
        control = next(n for n in timing.iter(f'{{{P}}}cTn') if n.get('nodeType') == 'interactiveSeq')
        control.find(f'{{{P}}}stCondLst/{{{P}}}cond/{{{P}}}tgtEl/{{{P}}}spTgt').set('spid', '9')
        ppt.save(self.source)
        issue = check_alignment(self.source, segments(1))[0]
        self.assertIsNone(issue.clicks)
        self.assertIn('指定对象', issue.reason)

    def test_media_player_without_a_main_sequence_has_zero_sequential_clicks(self):
        make_deck(self.source, [None])
        ppt = Presentation(self.source)
        ppt.slides[0]._element.append(parse_xml(
            f'<p:timing xmlns:p="{P}"><p:tnLst>{media_timing_xml(interactive=True)}</p:tnLst></p:timing>'))
        ppt.save(self.source)
        self.assertEqual(check_alignment(self.source, segments(1)), [])

    def test_document_order_is_used_instead_of_slide_part_filenames(self):
        make_deck(self.source, [None, [('clickEffect',)]])
        ppt = Presentation(self.source)
        ids = ppt.slides._sldIdLst
        ids.insert(0, ids[-1])
        ppt.save(self.source)
        self.assertEqual(check_alignment(self.source, segments(2, 1)), [])

    def test_unreferenced_slides_are_ignored_and_missing_page_is_reported(self):
        make_deck(self.source, [None, [('clickEffect',)]])
        self.assertEqual(check_alignment(self.source, segments(1)), [])
        issues = check_alignment(self.source, [{'page': 3, 'text': '无对应页'}])
        self.assertEqual(issues[0].page, 3)
        self.assertIn('超出', issues[0].reason)

    def test_independent_script_is_skipped_and_old_ppt_is_unconfirmed(self):
        self.assertEqual(check_alignment(None, segments(2)), [])
        self.assertIsNone(check_alignment(self.source.with_suffix('.ppt'), segments(2))[0].clicks)

    def test_blank_pages_are_reported_without_click_animations(self):
        make_deck(self.source, [None, None, None])
        current = [{'page': 1, 'text': '正常讲稿'},
                   {'page': 2, 'text': ' \n\t'},
                   {'page': 3, 'text': ''}]
        issues = check_alignment(self.source, current)
        self.assertEqual([(i.page, i.segments, i.clicks, i.reason) for i in issues],
                         [(page, 0, 0, '此页没有可朗读的讲稿') for page in (2, 3)])

    def test_blank_pages_keep_animation_findings_in_the_same_row(self):
        make_deck(self.source, [[('clickEffect',)], [('interactiveSeq',)]])
        issues = check_alignment(self.source, [{'page': 1, 'text': ''}, {'page': 2, 'text': ''}])
        self.assertEqual([(i.page, i.segments, i.clicks) for i in issues], [(1, 0, 1), (2, 0, None)])
        self.assertTrue(all('此页没有可朗读的讲稿' in i.reason for i in issues))
        self.assertIn('还有 1 个点击步骤', issues[0].reason)
        self.assertIn('点击指定对象', issues[1].reason)

    def test_empty_segments_warn_even_when_click_counts_match(self):
        make_deck(self.source, [[('clickEffect',)]])
        issues = check_alignment(self.source, [{'page': 1, 'text': ''}, {'page': 1, 'text': ' '}])
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0], AnimationIssue(1, 0, 1, '此页没有可朗读的讲稿'))


class AnimationGenerationTests(QtTestCase):

    def setUp(self):
        from ui.pages.reviewer_page import PPTReviewer
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.source = make_deck(Path(self.directory.name) / '讲稿.pptx', [None])
        ctx = NS(app_settings=AppSettings(), config=Mock(), playback_bus=PlaybackBus(),
                 hotkeys=Mock(), slideshow_watcher=Mock(), tts_engine=Mock())
        with patch('app.paths.ensure_runtime_dirs'):
            self.page = PPTReviewer(ctx)
        self.page.resize(1000, 800)
        self.page.project.presentation_path = self.source
        self.page.project.notes = {1: '第一段●第二段'}
        self.page.create_error_info_bar = Mock()
        self.addCleanup(self.dispose)

    def dispose(self):
        self.page.presentation_monitor.shutdown()
        self.page.stop_audio()
        self.page.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        self.qt.processEvents()

    def generate(self, decision):
        with patch.object(AnimationCheckDialog, 'exec', return_value=decision) as dialog, \
             patch.object(self.page.save_thread, 'configure') as configure, \
             patch.object(self.page.save_thread, 'start') as start, \
             patch.object(self.page, 'clean_temp_folder') as clean:
            result = self.page.regenerate()
        return result, dialog, configure, start, clean

    def test_cancel_restores_previous_segments_and_playable_audio(self):
        self.page.project.notes_list = [{'page': 1, 'text': '原有讲稿'}]
        audio = Path(self.directory.name) / 'old.wav'
        audio.write_bytes(b'keep cached audio on disk')
        self.page.project.media_list = [audio]
        self.page.project.audio_loaded = True
        result, dialog, configure, start, clean = self.generate(0)
        self.assertFalse(result)
        dialog.assert_called_once()
        configure.assert_not_called()
        start.assert_not_called()
        clean.assert_not_called()
        self.assertEqual(self.page.project.notes, {1: '第一段●第二段'})
        self.assertEqual(self.page.project.notes_list, [{'page': 1, 'text': '原有讲稿'}])
        self.assertEqual(self.page.project.media_list, [audio])
        self.assertTrue(audio.is_file())
        self.assertTrue(self.page.project.audio_loaded)
        self.assertTrue(self.page.getFileButton.isEnabled())
        self.assertFalse(self.page.is_busy())

    def test_continue_starts_generation_once_with_current_segments(self):
        result, dialog, configure, start, clean = self.generate(1)
        self.assertTrue(result)
        dialog.assert_called_once()
        start.assert_called_once_with()
        clean.assert_called_once()
        self.assertEqual(configure.call_args.args[0],
                         [{'page': 1, 'text': '第一段'}, {'page': 1, 'text': '第二段'}])

    def test_unchanged_retry_cancel_preserves_playable_audio(self):
        self.page.project.split_notes(comments_enabled=self.page.ctx.app_settings.get('script_comments_enabled'))
        self.page.project.media_list = [Path('a.wav'), Path('b.wav')]
        self.page.project.audio_loaded = True
        self.generate(0)
        self.assertEqual(self.page.project.media_list, [Path('a.wav'), Path('b.wav')])
        self.assertTrue(self.page.project.audio_loaded)

    def test_matched_page_and_standalone_script_generate_without_dialog(self):
        for source, text in ((self.source, '一段'), (None, '第一段●第二段')):
            self.page.project.presentation_path, self.page.project.notes = source, {1: text}
            result, dialog, configure, start, clean = self.generate(0)
            self.assertTrue(result)
            dialog.assert_not_called()
            start.assert_called_once_with()

    def test_disabled_check_skips_parsing_and_confirmation(self):
        self.page.ctx.app_settings.set('animation_check_enabled', False)
        self.source.write_bytes(b'not readable as pptx')
        with patch('app.presentation.animation.check_alignment') as check:
            result, dialog, configure, start, clean = self.generate(0)
        self.assertTrue(result)
        check.assert_not_called()
        dialog.assert_not_called()
        start.assert_called_once_with()

    def test_comments_filtered_before_generation_and_animation_check(self):
        make_deck(self.source, [[('clickEffect',)]])
        raw = '// 提醒●不参与分段\n第一段\n●第二段\n  // 结束●'
        self.page.ctx.app_settings.set('script_comments_enabled', True)
        self.page.project.notes = {1: raw}
        result, dialog, configure, start, clean = self.generate(0)
        self.assertTrue(result)
        dialog.assert_not_called()
        self.assertEqual(configure.call_args.args[0],
                         [{'page': 1, 'text': '第一段'}, {'page': 1, 'text': '第二段'}])
        self.assertEqual(self.page.project.notes, {1: raw})
        # 已有分段与统计保持生成时的规则，下次生成才应用新偏好。
        for enabled in (True, False):
            self.page.ctx.app_settings.set('script_comments_enabled', enabled)
            with patch('ui.pages.reviewer_page.ScriptStatisticsDialog') as statistics:
                self.page.show_info_dialog()
            self.assertEqual(statistics.call_args.args[0].characters, 6)
        self.page.project.split_notes(comments_enabled=self.page.ctx.app_settings.get('script_comments_enabled'))
        self.assertIn('// 提醒', self.page.project.notes_list[0]['text'])

    def test_comment_only_page_keeps_blank_page_timing_without_extra_clicks(self):
        self.page.ctx.app_settings.set('script_comments_enabled', True)
        self.page.project.notes = {1: '// 提示●\n\t// 资料●'}
        self.page.project.split_notes(comments_enabled=self.page.ctx.app_settings.get('script_comments_enabled'))
        self.assertEqual(self.page.project.notes_list, [{'page': 1, 'text': ''}])

    def test_blank_and_comment_only_pages_require_confirmation_before_generation(self):
        self.page.ctx.app_settings.set('script_comments_enabled', True)
        for text in (' \n\t', '// 提示●\n\t// 资料●'):
            for decision in (0, 1):
                with self.subTest(text=text, decision=decision):
                    self.page.project.notes = {1: text}
                    old_segments = list(self.page.project.notes_list)
                    result, dialog, configure, start, clean = self.generate(decision)
                    self.assertEqual(result, bool(decision))
                    dialog.assert_called_once()
                    self.assertEqual(self.page.project.notes_list,
                                     [{'page': 1, 'text': ''}] if decision else old_segments)
                    if decision:
                        start.assert_called_once_with()
                        configure.assert_called_once()
                    else:
                        start.assert_not_called()
                        configure.assert_not_called()

    def test_read_failure_still_allows_user_to_continue(self):
        self.source.write_bytes(b'broken ppt')
        result, dialog, configure, start, clean = self.generate(1)
        self.assertTrue(result)
        dialog.assert_called_once()
        start.assert_called_once_with()

    def test_current_multi_character_separator_and_trailing_mark_are_used(self):
        make_deck(self.source, [[('clickEffect',)]])
        self.page.project.mark = '||'
        self.page.project.notes = {1: '第一段||第二段||\n'}
        result, dialog, configure, start, clean = self.generate(0)
        self.assertTrue(result)
        dialog.assert_not_called()
        self.assertEqual(len(configure.call_args.args[0]), 2)

    def test_confirmation_prevents_reentrant_import_and_resets_busy_state(self):
        def confirm(dialog):
            self.assertTrue(self.page.is_busy())
            self.assertFalse(self.page.import_file(self.source))
            return 1
        self.page.create_warning_info_bar = Mock()
        with patch.object(AnimationCheckDialog, 'exec', confirm), \
             patch.object(self.page.save_thread, 'start') as start, \
             patch.object(self.page.save_thread, 'configure'), \
             patch.object(self.page, 'clean_temp_folder'):
            self.assertTrue(self.page.regenerate())
        start.assert_called_once_with()
        self.assertFalse(self.page.is_busy())

    def test_cancel_after_ppt_import_restores_empty_project_without_starting_generation(self):
        ppt = Presentation(self.source)
        ppt.slides[0].notes_slide.notes_text_frame.text = '导入第一段●导入第二段'
        ppt.save(self.source)
        self.page.project.notes = {}
        with patch.object(AnimationCheckDialog, 'exec', return_value=0), \
             patch.object(self.page.save_thread, 'start') as start:
            self.assertFalse(self.page.import_file(self.source))
        start.assert_not_called()
        self.assertEqual(self.page.project.notes, {})
        self.assertEqual(self.page.project.notes_list, [])
        self.assertFalse(self.page.project.audio_loaded)
        self.assertTrue(self.page.getFileButton.isEnabled())

    def test_many_issues_scroll_and_buttons_accept_or_cancel(self):
        issues = [AnimationIssue(i, 3, 1, '讲稿多出 1 个点击步骤') for i in range(1, 101)]
        dialog = AnimationCheckDialog(issues, self.page)
        dialog.show()
        self.qt.processEvents()
        self.assertEqual(dialog.table.rowCount(), 100)
        self.assertLessEqual(dialog.table.height(), 280)
        self.assertGreater(dialog.table.verticalScrollBar().maximum(), 0)
        self.assertEqual(dialog.yesButton.text(), '继续生成')
        self.assertEqual(dialog.cancelButton.text(), '取消')
        QTimer.singleShot(0, dialog.yesButton.click)
        self.assertEqual(dialog.exec(), 1)
        dialog = AnimationCheckDialog(issues[:1], self.page)
        QTimer.singleShot(0, dialog.cancelButton.click)
        self.assertEqual(dialog.exec(), 0)


if __name__ == '__main__':
    unittest.main()
