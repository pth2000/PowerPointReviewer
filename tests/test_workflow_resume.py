"""验证 PPT 更新需确认、后台检测及恢复上次讲稿时保持停止。"""

import json
from pathlib import Path
import tempfile
import wave
from types import SimpleNamespace as NS
import unittest
from unittest.mock import Mock, patch

from pptx import Presentation
from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, QTimer
from PySide6.QtMultimedia import QMediaPlayer
from PySide6.QtWidgets import QWidget

from app.presentation import link as presentation_link
from app.documents import io as script_io
from app.audio.playback import PlaybackBus
from app.presentation.monitor import changed_pages
from app.settings_store import AppSettings
from tasks.audio_generation_task import GenerationResult
from test_presentation_link import make_ppt
from ui.pages.reviewer_page import PPTReviewer

from qt_test_case import QtTestCase


class WorkflowResumeTests(QtTestCase):
    def start_update(self, notes):
        self.ctx.app_settings.set('animation_check_enabled', False)
        self.page.notes = dict(notes)
        with patch.object(self.page.save_thread, 'configure'), \
             patch.object(self.page.save_thread, 'start'), \
             patch.object(self.page, 'clean_temp_folder'):
            self.assertTrue(self.page.regenerate())

    def finish_update(self):
        count = len(self.page.notes_list)
        result = GenerationResult([self.page.audio_cache_path / 'key1.wav'] * count,
                                  [0.1] * count, ['key1'] * count, ['wav'] * count,
                                  generation_profile={'voice': 'original'}, speaker='测试声音')
        self.page.thread_save_finish(result)

    def test_update_returns_to_page_start_even_when_segment_count_is_unchanged(self):
        self.load()
        self.page.notes_list = [{'page': 1, 'text': '前页'}, {'page': 2, 'text': '第一段'},
                                {'page': 2, 'text': '第二段'}]
        self.page.current_index = 2
        self.start_update({1: '新增第一段●新增第二段●新增第三段', 2: '第一段●修改后的第二段'})
        self.assertEqual(self.page.current_index, 0)
        self.finish_update()
        self.assertEqual(self.page.current_index, 3)
        self.assertEqual(self.ctx.app_settings.get('last_session')['index'], 3)
        self.assertEqual(self.page.currentPageLabel.text(), '2 / 2')
        self.assertEqual(self.page._segment_offset(), 0)
        self.assertFalse(self.page._playing_sequence)
        self.assertEqual(self.page.player.playbackState(), QMediaPlayer.PlaybackState.StoppedState)

        self.page.scrollEnableSwitch.setChecked(True)
        with patch.object(self.page.playback_show, 'prepare') as prepare:
            self.page.init_play()
        self.assertEqual(prepare.call_args.args, (self.source, 2, 0))

    def test_changed_page_segmentation_returns_to_that_page_start(self):
        self.load()
        self.page.notes_list = [{'page': 1, 'text': '前页'}, {'page': 2, 'text': '第一段'},
                                {'page': 2, 'text': '第二段'}]
        self.page.current_index = 2
        self.start_update({1: '前页', 2: '合并后的讲稿'})
        self.finish_update()
        self.assertEqual(self.page.current_index, 1)
        self.assertEqual(self.page.notes_list[1]['page'], 2)

    def test_removed_page_falls_back_to_nearest_remaining_page(self):
        self.load()
        self.page.current_index = 1
        self.start_update({1: '保留的讲稿'})
        self.finish_update()
        self.assertEqual(self.page.current_index, 0)
        self.assertEqual(self.page.notes_list[0]['page'], 1)

    def test_repeated_updates_keep_one_record_and_preserve_its_identity(self):
        self.record.update(session_id='stable-id', created_at='2026-10-01T10:00:00', custom='keep')
        self.write_record()
        self.load()
        for text in ('第一轮修改', '第二轮修改'):
            self.start_update({1: self.baseline[1], 2: text})
            self.finish_update()
            self.assertEqual(self.page.session_record_path, self.record_path)
            self.assertEqual(list(self.root.glob('*.json')), [self.record_path])
            saved = json.loads(self.record_path.read_text(encoding='utf-8'))
            self.assertEqual(saved['session_id'], 'stable-id')
            self.assertEqual(saved['created_at'], '2026-10-01T10:00:00')
            self.assertEqual(saved['custom'], 'keep')
            self.assertEqual(saved['notes']['2'], text)
            self.assertIn('updated_at', saved)

    def test_failed_update_keeps_record_and_retry_restores_original_position(self):
        self.load()
        self.page.notes_list = [{'page': 1, 'text': '前页'}, {'page': 2, 'text': '第一段'},
                                {'page': 2, 'text': '第二段'}]
        self.page.current_index = 2
        original = self.record_path.read_bytes()
        self.ctx.app_settings.set('script_comments_enabled', True)
        self.start_update({1: '前页新增一段●前页新增二段', 2: '第二页更新第一段●第二页更新第二段'})
        with patch.object(self.page, 'create_error_info_bar'):
            self.page.thread_save_error('temporary synthesis failure')
        self.assertEqual(self.record_path.read_bytes(), original)
        self.assertEqual(self.page.session_record_path, self.record_path)
        annotations = script_io.ScriptData(notes={1: '// 新提示\n' + self.page.notes[1],
                                                 2: self.page.notes[2]})
        self.page._sync_ppt_comments((annotations, self.page.presentation_link))
        self.page._accept_ppt_baseline(self.page.presentation_link, annotations.notes)
        self.assertEqual(self.record_path.read_bytes(), original)
        self.start_update(self.page.notes)
        self.finish_update()
        self.assertEqual(self.page.current_index, 2)
        self.assertEqual(len(list(self.root.glob('*.json'))), 1)

    def test_atomic_record_write_failure_preserves_previous_record(self):
        self.load()
        original = self.record_path.read_bytes()
        with patch.object(Path, 'replace', side_effect=OSError('disk failure')):
            with self.assertRaises(OSError):
                self.page.save_session_record()
        self.assertEqual(self.record_path.read_bytes(), original)
        self.assertEqual(list(self.root.glob('*.tmp')), [])
        self.assertEqual(self.page.session_record_path, self.record_path)

    def test_importing_another_ppt_creates_a_separate_record(self):
        self.load()
        self.page.current_index = 1
        original = self.record_path.read_bytes()
        other = make_ppt(self.root / 'other.pptx')
        self.ctx.app_settings.set('animation_check_enabled', False)
        with patch.object(self.page.save_thread, 'configure'), \
             patch.object(self.page.save_thread, 'start'), \
             patch.object(self.page, 'clean_temp_folder'):
            self.assertTrue(self.page.import_file(other))
        self.assertIsNone(self.page.session_record_path)
        self.finish_update()
        self.assertNotEqual(self.page.session_record_path, self.record_path)
        self.assertEqual(self.page.current_index, 0)
        self.assertEqual(self.record_path.read_bytes(), original)
        self.assertEqual(len(list(self.root.glob('*.json'))), 2)

    def test_reimporting_current_ppt_retains_record_and_position(self):
        self.load()
        self.page.notes_list = [{'page': 1, 'text': '前页'}, {'page': 2, 'text': '第一段'},
                                {'page': 2, 'text': '第二段'}]
        self.page.current_index = 2
        self.edit_ppt('更新第一段●更新第二段')
        self.ctx.app_settings.set('animation_check_enabled', False)
        with patch.object(self.page.save_thread, 'configure'), \
             patch.object(self.page.save_thread, 'start'), \
             patch.object(self.page, 'clean_temp_folder'):
            self.assertTrue(self.page.import_file(self.source))
        self.assertEqual(self.page.session_record_path, self.record_path)
        self.finish_update()
        self.assertEqual(self.page.current_index, 1)
        self.assertEqual(len(list(self.root.glob('*.json'))), 1)

    def test_history_shows_update_time_and_supports_older_records(self):
        from ui.dialogs.session_history_dialog import SessionHistoryDialog
        self.record.update(created_at='2026-10-01T10:00:00', updated_at='2026-10-18T12:00:00')
        self.write_record()
        with patch('app.audio.cache.summary', return_value='0 条'):
            dialog = SessionHistoryDialog(self.root, self.page)
        self.addCleanup(dialog.deleteLater)
        self.assertEqual(dialog.table.horizontalHeaderItem(0).text(), '更新时间')
        self.assertEqual(dialog.table.item(0, 0).text(), '2026-10-18 12:00:00')
        self.record.pop('updated_at')
        self.write_record()
        dialog.reload_records()
        self.assertEqual(dialog.table.item(0, 0).text(), '2026-10-01 10:00:00')

    def test_comment_history_and_package_keep_generation_rule_and_raw_notes(self):
        from app.documents import package as project_package
        self.record['script_comments_enabled'] = True
        self.record['notes'] = {int(page): '// 提示●\n' + text for page, text in self.baseline.items()}
        self.write_record()
        self.load()
        self.assertFalse(self.ctx.app_settings.get('script_comments_enabled'))
        self.assertTrue(self.page.script_comments_enabled)
        self.assertNotIn('//', ''.join(self.page.reading_notes().values()))
        self.assertIn('// 提示', self.page.notes[1])
        self.page.save_session_record()
        saved = json.loads(self.page.session_record_path.read_text(encoding='utf-8'))
        self.assertTrue(saved['script_comments_enabled'])
        self.assertEqual(saved['notes'], {str(page): text for page, text in self.page.notes.items()})
        target = self.root / 'comments.pprpkg'
        project_package.export_package(target, items=self.record['items'], notes=self.page.notes,
                                       mark='●', source_name='讲稿', generation_profile={}, speaker='',
                                       include_audio=False, script_comments_enabled=True)
        manifest = project_package.read_manifest(target)
        self.assertTrue(manifest['script_comments_enabled'])
        self.assertIn('// 提示', manifest['notes']['1'])
        with patch('app.paths.ensure_runtime_dirs'), patch('app.paths.SESSION_DIR', self.root):
            result = project_package.import_package(target)
        imported = json.loads(result['record_path'].read_text(encoding='utf-8'))
        self.assertTrue(imported['script_comments_enabled'])

    def test_monitor_syncs_comments_without_prompting_or_replacing_ai_body(self):
        self.record['script_comments_enabled'] = True
        self.write_record()
        self.load()
        self.page.notes[2] = 'AI 改写后的讲稿'
        old_media = list(self.page.media_list)
        self.edit_ppt('// 新的教学提示●\n' + self.baseline[2])
        with patch.object(self.page, 'regenerate') as generate:
            self.check_monitor()
        self.assertIsNone(self.page._ppt_update_bar)
        generate.assert_not_called()
        self.assertEqual(self.page.media_list, old_media)
        self.assertEqual(self.page.notes[2], '// 新的教学提示●\nAI 改写后的讲稿')
        saved = json.loads(self.record_path.read_text(encoding='utf-8'))
        self.assertEqual(saved['items'], self.record['items'])
        self.assertEqual(saved['notes']['2'], self.page.notes[2])

    def test_ai_rewrite_preserves_comments_in_current_notes(self):
        self.record['script_comments_enabled'] = True
        self.record['notes'] = {**self.baseline, 1: '// 保留教学提示\n原文\n// 保留资料'}
        self.write_record()
        self.load()
        with patch.object(self.page, 'regenerate') as generate:
            self.assertEqual(self.page.apply_rewritten_notes({1: '改写后的正文'}), 1)
        self.assertEqual(self.page.reading_notes()[1], '改写后的正文')
        self.assertIn('// 保留教学提示', self.page.notes[1])
        self.assertIn('// 保留资料', self.page.notes[1])
        generate.assert_called_once_with()


    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.root = Path(folder.name)
        self.source = make_ppt(self.root / 'lesson.pptx')
        self.baseline = script_io.load_script(self.source).notes
        audio_dir = self.root / 'audio'
        audio_dir.mkdir()
        items = []
        for page, text in self.baseline.items():
            with wave.open(str(audio_dir / f'key{page}.wav'), 'wb') as audio:
                audio.setnchannels(1)
                audio.setsampwidth(2)
                audio.setframerate(24000)
                audio.writeframes(bytes(4800))
            items.append({'page': page, 'text': text.strip(), 'cache_key': f'key{page}',
                          'cache_ext': 'wav', 'duration': 1})
        self.record = {'source_file': str(self.source), 'source_name': '测试讲稿',
                       'presentation': presentation_link.describe(self.source),
                       'presentation_notes': self.baseline, 'notes': self.baseline,
                       'items': items, 'mark': '●', 'generation_profile': {'voice': 'original'}}
        self.record_path = self.root / 'history.json'
        self.write_record()
        self.ctx = NS(app_settings=AppSettings(), config=Mock(), playback_bus=PlaybackBus(),
                      slideshow_watcher=Mock(), hotkeys=Mock(), tts_engine=Mock())
        self.ctx.tts_engine.get_selected_voice_name.return_value = '测试声音'
        self.ctx.tts_engine.snapshot.return_value = self.ctx.tts_engine
        with patch('app.paths.ensure_runtime_dirs'):
            self.page = PPTReviewer(self.ctx)
        self.page.audio_cache_path = audio_dir
        self.page.session_root_path = self.root
        self.page.presentation_monitor.timer.stop()
        self.addCleanup(self.dispose)

    def dispose(self):
        self.page.presentation_monitor.shutdown()
        self.page.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        self.qt.processEvents()

    def write_record(self):
        self.record_path.write_text(json.dumps(self.record), encoding='utf-8')

    def load(self):
        self.page.load_session_record(self.record_path)

    def edit_ppt(self, text='新的第二页备注'):
        ppt = Presentation(str(self.source))
        ppt.slides[1].notes_slide.notes_text_frame.text = text
        ppt.save(str(self.source))

    def check_monitor(self):
        monitor = self.page.presentation_monitor
        monitor.poll()
        if not monitor.task.isRunning():
            self.qt.processEvents()
            return
        loop = QEventLoop()
        monitor.task.finished.connect(loop.quit)
        QTimer.singleShot(3000, loop.quit)
        loop.exec()
        monitor.task.finished.disconnect(loop.quit)
        self.assertFalse(monitor.task.isRunning())
        self.qt.processEvents()

    def wait_until(self, condition):
        if condition():
            return
        loop = QEventLoop()
        timer = QTimer()
        timer.setInterval(20)
        timer.timeout.connect(lambda: loop.quit() if condition() else None)
        timeout = QTimer()
        timeout.setSingleShot(True)
        timeout.timeout.connect(loop.quit)
        timer.start()
        timeout.start(4000)
        loop.exec()
        timer.stop()
        timeout.stop()
        self.assertTrue(condition())

    def test_filesystem_events_handle_atomic_save_delete_and_restore_without_polling(self):
        self.load()
        monitor = self.page.presentation_monitor
        self.wait_until(lambda: monitor.stamp is not None and not monitor.task.isRunning()
                        and not monitor.timer.isActive())
        self.assertTrue(monitor.timer.isSingleShot())
        replacement = self.root / 'replacement.pptx'
        replacement.write_bytes(self.source.read_bytes())
        ppt = Presentation(str(replacement))
        ppt.slides[1].notes_slide.notes_text_frame.text = '原子保存后的备注'
        ppt.save(str(replacement))
        replacement.replace(self.source)
        self.wait_until(lambda: self.page._ppt_update_bar is not None)
        payload = self.source.read_bytes()
        self.source.unlink()
        self.wait_until(lambda: self.page.is_restricted() and
                        not self.page.scrollEnableSwitch.isEnabled())
        self.source.write_bytes(payload)
        self.wait_until(lambda: not self.page.is_restricted() and
                        str(self.source) in monitor.watcher.files())
        self.wait_until(lambda: not monitor.task.isRunning() and not monitor.timer.isActive())
        self.assertFalse(self.page.device_watcher._timer.isActive())

    def test_new_dialog_shows_page_summary_without_script_text(self):
        from qfluentwidgets import BodyLabel, FluentWindow, FluentIcon
        from ui.dialogs.ppt_update_dialog import PptUpdateDialog, page_ranges
        root = FluentWindow()
        root.resize(760, 560)
        root.addSubInterface(self.page, FluentIcon.HOME, '主页')
        root.show()
        self.addCleanup(root.deleteLater)
        before = {1: '保持不变', 2: '旧正文不能显示', 3: '要删除的正文'}
        after = {1: '保持不变', 2: '新正文不能显示', 4: '新增正文不能显示'}
        dialog = PptUpdateDialog('example.pptx', before, after, True,
                                 'zh-CN-YunyangNeural', self.page)
        self.addCleanup(dialog.deleteLater)
        dialog.show()
        self.qt.processEvents()
        self.assertEqual(dialog.summary_label.text(), '3 页讲稿有变化')
        self.assertEqual([label.text() for label in dialog.page_labels],
                         ['第 2 页', '第 4 页', '第 3 页'])
        labels = '\n'.join(label.text() for label in dialog.findChildren(BodyLabel))
        self.assertNotIn('正文不能显示', labels)
        self.assertNotIn('要删除的正文', labels)
        self.assertLessEqual(dialog.widget.height(), self.page.height())
        self.assertTrue(dialog.widget.rect().contains(
            dialog.yesButton.mapTo(dialog.widget, dialog.yesButton.rect().bottomRight())))
        self.assertEqual(page_ranges([1, 2, 3, 7, 9, 10]), '1–3、7、9–10')
        dialog.reject()

    def test_background_detects_only_note_changes_without_replacing_audio(self):
        self.load()
        old_notes, old_media = dict(self.page.notes), list(self.page.media_list)
        with patch.object(self.page, '_notify_ppt_update'):
            self.check_monitor()
            self.assertIsNone(self.page._ppt_update_bar)
            # 只改标题、图片等画面内容不会提示备注更新。
            ppt = Presentation(str(self.source))
            ppt.core_properties.title = 'new title'
            ppt.save(str(self.source))
            self.check_monitor()
            self.assertIsNone(self.page._ppt_update_bar)
        self.edit_ppt()
        with patch.object(self.page, 'regenerate') as generate:
            self.check_monitor()
            self.assertIsNotNone(self.page._ppt_update_bar)
            generate.assert_not_called()
        self.assertEqual(self.page.notes, old_notes)
        self.assertEqual(self.page.media_list, old_media)
        self.assertEqual(self.page.current_index, 0)
        bar = self.page._ppt_update_bar
        self.check_monitor()
        self.assertIs(self.page._ppt_update_bar, bar)

    def test_cancel_update_preserves_ai_notes_audio_and_record(self):
        self.load()
        self.page.notes[1] = 'AI 改写后的第一段'
        old_record = self.record_path.read_bytes()
        old_notes = dict(self.page.notes)
        self.edit_ppt()
        with patch('ui.dialogs.ppt_update_dialog.PptUpdateDialog') as dialog, \
             patch.object(self.page, 'regenerate') as generate:
            dialog.return_value.exec.return_value = False
            self.page.update_ppt_notes()
            self.assertTrue(dialog.call_args.args[3])
            generate.assert_not_called()
        self.assertEqual(self.page.notes, old_notes)
        self.assertEqual(self.record_path.read_bytes(), old_record)
        self.assertTrue(self.page.is_import)

    def test_confirm_update_keeps_record_until_generation_finishes(self):
        self.load()
        original = self.record_path.read_bytes()
        self.edit_ppt()
        with patch('ui.dialogs.ppt_update_dialog.PptUpdateDialog') as dialog, \
             patch.object(self.page, 'regenerate') as generate:
            dialog.return_value.exec.return_value = True
            self.page.update_ppt_notes()
            generate.assert_called_once_with()
        self.assertEqual(self.page.notes[2].strip(), '新的第二页备注')
        self.assertEqual(self.page.session_record_path, self.record_path)
        self.assertEqual(self.record_path.read_bytes(), original)
        self.assertEqual(self.page.presentation_monitor.baseline, self.page.notes)

    def test_actual_update_generation_reuses_unchanged_audio(self):
        self.load()
        self.page.current_index = 1
        engine = self.ctx.tts_engine
        engine.get_generation_profile.return_value = {'voice': 'original'}
        engine.get_output_extension.return_value = 'wav'
        engine.can_parallel_generate.return_value = False
        engine.build_audio_cache_key.side_effect = lambda text, profile: (
            'key1' if text == self.baseline[1].strip() else
            'key2' if text == self.baseline[2].strip() else 'new-key')
        def synthesize(text, path):
            with wave.open(path, 'wb') as audio:
                audio.setnchannels(1)
                audio.setsampwidth(2)
                audio.setframerate(24000)
                audio.writeframes(bytes(4800))
        engine.save_file.side_effect = synthesize
        task = self.page.save_thread
        task.audio_cache_path = self.page.audio_cache_path
        task.output_path = self.root / 'temp'
        task.output_path.mkdir()
        self.page.wav_temp_path = task.output_path
        old_audio = (self.page.audio_cache_path / 'key1.wav').read_bytes()
        old_record = self.record_path.read_bytes()
        self.edit_ppt()
        with patch('ui.dialogs.ppt_update_dialog.PptUpdateDialog') as dialog, \
             patch.object(task, 'save_countdown_wav'):
            dialog.return_value.exec.return_value = True
            self.page.update_ppt_notes()
            loop = QEventLoop()
            task.finished.connect(loop.quit)
            QTimer.singleShot(3000, loop.quit)
            loop.exec()
            task.finished.disconnect(loop.quit)
            self.qt.processEvents()
        self.assertFalse(task.isRunning())
        self.assertTrue(self.page.is_import)
        self.assertEqual(self.page.cache_hit_count, 1)
        self.assertEqual(engine.save_file.call_count, 1)
        self.assertEqual(engine.save_file.call_args.args[0], '新的第二页备注')
        self.assertEqual((self.page.audio_cache_path / 'key1.wav').read_bytes(), old_audio)
        self.assertNotEqual(self.record_path.read_bytes(), old_record)
        self.assertEqual(self.page.session_record_path, self.record_path)
        self.assertEqual(len(list(self.root.glob('*.json'))), 1)
        updated = json.loads(self.record_path.read_text(encoding='utf-8'))
        self.assertEqual(updated['notes']['2'].strip(), '新的第二页备注')
        self.assertEqual(self.page.current_index, 1)
        self.assertFalse(self.page._playing_sequence)
        self.assertEqual(self.page.player.playbackState(), QMediaPlayer.PlaybackState.StoppedState)
        self.assertEqual(self.ctx.app_settings.get('last_session')['record_path'],
                         str(self.page.session_record_path.resolve()))

    def test_save_while_confirmation_is_open_requires_check_again(self):
        self.load()
        original = dict(self.page.notes)
        self.edit_ppt()
        with patch('ui.dialogs.ppt_update_dialog.PptUpdateDialog') as dialog, \
             patch.object(self.page, 'regenerate') as generate, \
             patch.object(self.page, 'create_error_info_bar') as error:
            def confirm():
                self.edit_ppt('确认期间再次编辑')
                return True
            dialog.return_value.exec.side_effect = confirm
            self.page.update_ppt_notes()
            error.assert_called_once()
            generate.assert_not_called()
        self.assertEqual(self.page.notes, original)

    def test_same_notes_do_not_regenerate_and_read_failure_preserves_state(self):
        self.load()
        with patch.object(self.page, 'regenerate') as generate:
            self.page.update_ppt_notes()
            generate.assert_not_called()
        original = dict(self.page.notes)
        self.source.write_bytes(b'partial save')
        with patch.object(self.page, 'create_error_info_bar') as error:
            self.page.update_ppt_notes()
            error.assert_called_once()
        self.assertEqual(self.page.notes, original)
        self.assertTrue(self.page.is_import)

    def test_changed_ppt_since_last_run_prompts_without_confusing_ai_edits(self):
        self.record['notes'] = {1: 'AI 改写', 2: self.baseline[2]}
        self.write_record()
        self.load()
        self.check_monitor()
        self.assertIsNone(self.page._ppt_update_bar)
        self.edit_ppt()
        self.load()
        self.check_monitor()
        self.assertIsNotNone(self.page._ppt_update_bar)

    def test_new_association_sets_new_baseline_and_ignores_stale_result(self):
        self.load()
        monitor = self.page.presentation_monitor
        old_epoch = monitor.epoch
        other = make_ppt(self.root / 'other.pptx')
        self.page.set_presentation_link(presentation_link.describe(other))
        self.check_monitor()
        self.assertIsNone(self.page._ppt_update_bar)
        monitor._checked((old_epoch, (0, 0), NS(notes={1: 'stale'}), {'sha256': 'stale'}))
        self.assertEqual(monitor.baseline, script_io.load_script(other).notes)
        self.assertIsNone(self.page._ppt_update_bar)
        record = json.loads(self.record_path.read_text(encoding='utf-8'))
        self.assertIsNone(record['presentation_notes'])

    def test_reverted_ppt_dismisses_update_and_busy_check_is_deferred(self):
        self.load()
        self.edit_ppt()
        self.check_monitor()
        self.assertIsNotNone(self.page._ppt_update_bar)
        monitor = self.page.presentation_monitor
        previous_stamp = monitor.stamp
        with patch.object(self.page, 'is_busy', return_value=True):
            monitor._checked((monitor.epoch, (1, 1), NS(notes={1: 'later'}), {'sha256': 'later'}))
        self.assertEqual(monitor.stamp, previous_stamp)
        self.edit_ppt(self.baseline[2].strip())
        self.check_monitor()
        self.assertIsNone(self.page._ppt_update_bar)

    def test_background_read_failure_retries_after_file_is_saved(self):
        self.load()
        original = self.source.read_bytes()
        self.source.write_bytes(b'incomplete ppt')
        self.check_monitor()
        self.assertIsNone(self.page._ppt_update_bar)
        self.assertIsNone(self.page.presentation_monitor.stamp)
        self.source.write_bytes(original)
        self.edit_ppt()
        self.check_monitor()
        self.assertIsNotNone(self.page._ppt_update_bar)

    def test_resume_restores_segment_stopped_without_generating_or_show_control(self):
        self.ctx.app_settings.set('last_session', {'record_path': str(self.record_path), 'index': 1})
        with patch('ui.pages.reviewer_page.MessageBox') as box, \
             patch.object(self.page.save_thread, 'start') as generate, \
             patch('app.presentation.slideshow.goto') as go:
            box.return_value.exec.return_value = True
            self.page.prompt_resume_session()
            self.assertIn('第 2 页', box.call_args.args[1])
            generate.assert_not_called()
            go.assert_not_called()
        self.assertEqual(self.page.current_index, 1)
        self.assertEqual(self.page.player.playbackState(), QMediaPlayer.PlaybackState.StoppedState)
        self.assertEqual(self.ctx.app_settings.get('last_session')['index'], 1)

    def window_with_hidden_home(self):
        from qfluentwidgets import FluentWindow, FluentIcon
        root = FluentWindow()
        root.resize(850, 750)
        root.addSubInterface(self.page, FluentIcon.HOME, '主页')
        other = QWidget(root)
        other.setObjectName('settings-test')
        root.addSubInterface(other, FluentIcon.SETTING, '设置')
        self.addCleanup(root.deleteLater)
        root.show()
        root.stackedWidget.setCurrentWidget(other)
        self.wait_until(lambda: not self.page.isVisible())
        return root

    def test_resume_from_settings_uses_visible_main_window_as_parent(self):
        root = self.window_with_hidden_home()
        self.ctx.app_settings.set('last_session', {'record_path': str(self.record_path)})
        with patch('ui.pages.reviewer_page.MessageBox') as box:
            box.return_value.exec.return_value = False
            self.page.prompt_resume_session()
        self.assertIs(box.call_args.args[2], root)
        self.assertTrue(root.isVisible())

    def test_resume_waits_for_embedded_color_dialog_and_only_prompts_once(self):
        from PySide6.QtGui import QColor
        from ui import theme
        root = self.window_with_hidden_home()
        self.ctx.app_settings.set('last_session', {'record_path': str(self.record_path)})
        color = theme.create_color_dialog(QColor('#B7472A'), root)
        self.addCleanup(color.deleteLater)
        color.setModal(True)
        color.show()
        self.qt.processEvents()
        # Fluent 遮罩弹窗并非独立窗口，Qt 的 activeModalWidget 检测不到它。
        self.assertFalse(color.isWindow())
        self.assertIsNone(self.qt.activeModalWidget())
        with patch('ui.pages.reviewer_page.MessageBox') as box:
            box.return_value.exec.return_value = False
            self.page.prompt_resume_session()
            self.page.prompt_resume_session()
            box.assert_not_called()
            self.assertTrue(color.isVisible())
            color.reject()
            self.wait_until(lambda: box.call_count == 1)
            self.assertIs(box.call_args.args[2], root)
            self.assertFalse(self.page._resume_waiting_for_dialog)

    def test_deferred_resume_preserves_script_imported_while_waiting(self):
        from PySide6.QtGui import QColor
        from ui import theme
        root = self.window_with_hidden_home()
        self.ctx.app_settings.set('last_session', {'record_path': str(self.record_path)})
        color = theme.create_color_dialog(QColor('#B7472A'), root)
        self.addCleanup(color.deleteLater)
        color.setModal(True)
        color.show()
        self.qt.processEvents()
        with patch('ui.pages.reviewer_page.MessageBox') as box:
            self.page.prompt_resume_session()
            self.page.notes = {1: '当前讲稿'}
            color.reject()
            self.wait_until(lambda: not self.page._resume_waiting_for_dialog)
            self.qt.processEvents()
            box.assert_not_called()
        self.assertEqual(self.page.notes, {1: '当前讲稿'})

    def test_disabled_resume_skips_prompt_and_keeps_saved_session(self):
        state = {'record_path': str(self.record_path), 'index': 1}
        self.ctx.app_settings.set('last_session', state)
        self.ctx.app_settings.set('resume_last_session', False)
        with patch('ui.pages.reviewer_page.MessageBox') as box, \
             patch.object(self.page, 'load_session_record') as load, \
             patch.object(Path, 'read_text') as read:
            self.page.prompt_resume_session()
            box.assert_not_called()
            load.assert_not_called()
            read.assert_not_called()
        self.assertEqual(self.ctx.app_settings.get('last_session'), state)
        self.assertFalse(self.page.notes)

    def test_resume_cancel_missing_record_and_existing_work_do_not_load(self):
        self.ctx.app_settings.set('last_session', {'record_path': str(self.record_path), 'index': 0})
        with patch('ui.pages.reviewer_page.MessageBox') as box, \
             patch.object(self.page, 'load_session_record') as load:
            box.return_value.exec.return_value = False
            self.page.prompt_resume_session()
            load.assert_not_called()
            box.reset_mock()
            self.page.notes = {1: '用户已经导入的讲稿'}
            self.page.prompt_resume_session()
            box.assert_not_called()
            self.page.notes = {}
            self.record_path.unlink()
            self.page.prompt_resume_session()
            box.assert_not_called()

    def test_resume_clamps_index_and_deleted_ppt_loads_restricted_audio(self):
        self.ctx.app_settings.set('last_session', {'record_path': str(self.record_path), 'index': 99})
        payload = self.source.read_bytes()
        self.source.unlink()
        with patch('ui.pages.reviewer_page.MessageBox') as box, \
             patch('app.presentation.slideshow.list_open_presentations', return_value=[]):
            box.return_value.exec.return_value = True
            self.page.prompt_resume_session()
        self.assertEqual(self.page.current_index, 1)
        self.assertTrue(self.page.is_restricted())
        self.assertTrue(self.page.playButton.isEnabled())
        self.assertFalse(self.page.scrollEnableSwitch.isEnabled())
        self.source.write_bytes(payload)
        self.wait_until(lambda: not self.page.is_restricted())
        self.assertEqual(self.page.presentation_path, self.source)

    def test_failed_generation_does_not_replace_last_successful_session(self):
        self.load()
        self.page.current_index = 1
        self.page.set_current_label_text()
        saved = dict(self.ctx.app_settings.get('last_session'))
        self.page.clean_and_reset()
        self.page.remember_session()
        self.assertEqual(self.ctx.app_settings.get('last_session'), saved)

    def test_baseline_survives_new_history_and_whitespace_is_not_a_change(self):
        self.load()
        self.ctx.tts_engine.get_generation_profile.return_value = {'voice': 'test'}
        self.page.save_session_record()
        record = json.loads(self.page.session_record_path.read_text(encoding='utf-8'))
        self.assertEqual(changed_pages(record['presentation_notes'], self.baseline), [])
        self.assertEqual(changed_pages({1: '备注\r\n'}, {1: '备注\n'}), [])


if __name__ == '__main__':
    unittest.main()
